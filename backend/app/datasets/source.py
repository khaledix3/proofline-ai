"""Read-only UCI workbook adapter. Invoice lines are not unique invoices/orders."""
from dataclasses import dataclass
from datetime import datetime
from decimal import Decimal
from hashlib import sha256
from pathlib import Path
from openpyxl import load_workbook
from app.replay.inputs import fingerprint

DATASET_URL = "https://archive.ics.uci.edu/dataset/352/online+retail"
DOWNLOAD_URL = "https://archive.ics.uci.edu/static/public/352/online%2Bretail.zip"
CITATION = "Chen, D. (2015). Online Retail. UCI Machine Learning Repository. https://doi.org/10.24432/C5BW33"
OFFICIAL_WORKBOOK_SHA256 = "43465a06f2ccf7c8b5bd2892bc7defb52f97487934fe93b16ae4c3936424676d"
FIELDS = ("InvoiceNo", "StockCode", "Description", "Quantity", "InvoiceDate", "UnitPrice", "CustomerID", "Country")
MAPPING_VERSION = "uci352.invoice-line.v1"

@dataclass(frozen=True)
class RetailSubset:
    workbook_sha256: str
    sheet: str
    rows: tuple[dict, ...]
    selection: str

    @property
    def digest(self):
        return fingerprint({"sheet": self.sheet, "rows": self.rows})


def read_subset(path: Path, size: int = 1000) -> RetailSubset:
    if not 500 <= size <= 2000:
        raise ValueError("Select between 500 and 2000 source rows")
    digest = sha256(path.read_bytes()).hexdigest()
    if digest != OFFICIAL_WORKBOOK_SHA256:
        raise ValueError("Workbook does not match the acquired official UCI artifact; reverify provenance")
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.worksheets[0]
        iterator = sheet.iter_rows(values_only=True)
        if tuple(next(iterator)) != FIELDS:
            raise ValueError("Unexpected source workbook fields")
        total = sheet.max_row - 1
        if total < size:
            raise ValueError("Workbook is smaller than the requested subset")
        targets = {2 + i * (total - 1) // (size - 1) for i in range(size)}
        rows = []
        for number, values in enumerate(iterator, 2):
            if number in targets:
                original = {key: value.isoformat() if isinstance(value, datetime) else value for key, value in zip(FIELDS, values)}
                rows.append({"source_row": number, "original": original})
        if len(rows) != size:
            raise ValueError("Workbook dimensions did not match readable rows")
        return RetailSubset(digest, sheet.title, tuple(rows), f"{size} evenly spaced worksheet data rows, including first/last; original row order")
    finally:
        workbook.close()


def source_identity(subset: RetailSubset, row: dict) -> str:
    return f"{subset.workbook_sha256}:{row['source_row']}"


def cell_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float):
        return format(Decimal(str(value)), "f")
    return str(value)


def map_line(subset: RetailSubset, row: dict) -> dict[str, str]:
    original = row["original"]
    customer = cell_text(original.get("CustomerID"))
    invoice_date = cell_text(original.get("InvoiceDate"))
    # This is a technical source-line identity, NEVER a claim that InvoiceNo is unique.
    return {"order_id": f"UCI352_{subset.workbook_sha256[:12]}_{row['source_row']}",
            "customer_id": "UCI_C_" + sha256(customer.encode()).hexdigest()[:16] if customer else "",
            "quantity": cell_text(original.get("Quantity")),
            "unit_price": cell_text(original.get("UnitPrice")),
            "order_date": invoice_date.split("T")[0]}


def classify_line(original: dict) -> str:
    # Descriptive primary categories, not fraud labels or detector outcomes.
    if cell_text(original.get("InvoiceNo")).upper().startswith("C"):
        return "cancellation"
    quantity = original.get("Quantity")
    if isinstance(quantity, (int, float)) and quantity < 0:
        return "negative_quantity_return_or_adjustment_candidate"
    if quantity == 0:
        return "zero_quantity"
    if original.get("CustomerID") is None:
        return "missing_customer_reference"
    if original.get("UnitPrice") == 0:
        return "zero_price"
    return "other_transaction_line"
